import io
import pytest
from fastapi.testclient import TestClient
from sqlalchemy.orm import Session
import openpyxl

from app.models.client import Client
from app.models.project import Project
from app.models.survey import Survey, SurveyQuestion, SurveyAnswer
from app.services.export_service import (
    export_project_analytics,
    export_surveys_registry,
    export_portfolio_overview,
)

def test_export_project_analytics_csv_and_xlsx(db: Session, test_user):
    # Setup test client & project
    client = Client(client_code="CLI-EXP-01", name="Export Test Client", is_active=True)
    db.add(client)
    db.flush()

    project = Project(
        name="Export Test Project",
        project_code="EXP-001",
        client_id=client.id,
        team_name="Export Lead",
        is_active=True
    )
    db.add(project)
    db.flush()

    survey = Survey(
        project_id=project.id,
        client_id=client.id,
        creator_id=test_user.id,
        phase="Mid-Project",
        status="Completed",
        survey_code="EXP-001-MID-01",
        unique_token="exp-test-token-001",
        average_score=4.5,
        client_contact_name="Client Rep",
        client_contact_email="rep@client.com",
    )
    db.add(survey)
    db.flush()

    q1 = SurveyQuestion(survey_id=survey.id, section_number=1, section_title="Project Management", question_text="PM Communication", response_type="star_rating")
    q6 = SurveyQuestion(survey_id=survey.id, section_number=6, section_title="Qualitative Feedback", question_text="Comments", response_type="free_text")
    db.add_all([q1, q6])
    db.flush()

    a1 = SurveyAnswer(survey_id=survey.id, question_id=q1.id, rating_score=5)
    a6 = SurveyAnswer(survey_id=survey.id, question_id=q6.id, text_answer="Great collaboration throughout the sprint.")
    db.add_all([a1, a6])
    db.commit()

    # 1. Test CSV export
    csv_bytes, csv_filename, csv_mime = export_project_analytics(db, project.id, format="csv")
    assert csv_mime == "text/csv"
    assert "EXP-001_CSAT_Report.csv" in csv_filename
    csv_text = csv_bytes.decode("utf-8")
    assert "Export Test Project" in csv_text
    assert "EXP-001" in csv_text
    assert "Great collaboration throughout the sprint." in csv_text

    # 2. Test XLSX export
    xlsx_bytes, xlsx_filename, xlsx_mime = export_project_analytics(db, project.id, format="xlsx")
    assert "openxmlformats" in xlsx_mime
    assert "EXP-001_CSAT_Report.xlsx" in xlsx_filename
    
    # Parse workbook from in-memory bytes
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    sheet_names = wb.sheetnames
    assert "Project CSAT Summary" in sheet_names
    assert "Survey Registry" in sheet_names
    assert "Qualitative Feedback" in sheet_names

    ws_summary = wb["Project CSAT Summary"]
    assert "EXP-001" in str(ws_summary["A1"].value)
    ws_comments = wb["Qualitative Feedback"]
    found_comment = any("Great collaboration" in str(cell.value) for row in ws_comments.rows for cell in row)
    assert found_comment

def test_export_surveys_registry_and_portfolio_overview(db: Session):
    # Test surveys registry export
    csv_bytes, filename, mime = export_surveys_registry(db, format="csv")
    assert mime == "text/csv"
    assert filename == "CSAT_Surveys_Registry.csv"
    assert len(csv_bytes) > 0

    xlsx_bytes, filename, mime = export_surveys_registry(db, format="xlsx")
    assert "openxmlformats" in mime
    assert filename == "CSAT_Surveys_Registry.xlsx"
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    assert "CSAT Surveys" in wb.sheetnames

    # Test portfolio overview export
    csv_bytes, filename, mime = export_portfolio_overview(db, format="csv")
    assert mime == "text/csv"
    assert filename == "CSAT_Portfolio_Overview.csv"

    xlsx_bytes, filename, mime = export_portfolio_overview(db, format="xlsx")
    assert "openxmlformats" in mime
    assert filename == "CSAT_Portfolio_Overview.xlsx"
    wb_portfolio = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    assert "Executive CSAT Summary" in wb_portfolio.sheetnames

def test_api_export_endpoints(client: TestClient, auth_headers: dict, db: Session):
    # Setup a project
    c = Client(client_code="CLI-API-01", name="API Test Client", is_active=True)
    db.add(c)
    db.flush()
    p = Project(name="API Test Project", project_code="API-001", client_id=c.id, is_active=True)
    db.add(p)
    db.commit()

    # 1. Test Project Export Endpoint CSV
    res_csv = client.get(f"/api/analytics/projects/{p.id}/export?format=csv", headers=auth_headers)
    assert res_csv.status_code == 200
    assert res_csv.headers["content-type"] == "text/csv; charset=utf-8"
    assert "attachment" in res_csv.headers["content-disposition"]

    # 2. Test Project Export Endpoint XLSX
    res_xlsx = client.get(f"/api/analytics/projects/{p.id}/export?format=xlsx", headers=auth_headers)
    assert res_xlsx.status_code == 200
    assert "application/vnd.openxmlformats" in res_xlsx.headers["content-type"]

    # 3. Test Surveys Registry Export Endpoint
    res_surveys = client.get("/api/analytics/surveys/export?format=csv", headers=auth_headers)
    assert res_surveys.status_code == 200

    # 4. Test Portfolio Overview Export Endpoint
    res_overview = client.get("/api/analytics/overview/export?format=xlsx", headers=auth_headers)
    assert res_overview.status_code == 200

    # 5. Test Email Service Status Endpoint
    res_email = client.get("/api/surveys/email-status", headers=auth_headers)
    assert res_email.status_code == 200
    email_data = res_email.json()
    assert "configured" in email_data
    assert "host" in email_data
